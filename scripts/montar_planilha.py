import json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

d = json.load(open(sys.argv[1], encoding="utf-8"))
out = sys.argv[2]
HOJE = "08/10/2026"
ORIGEM = "Encaminha Goiânia, index.html (commit aead8b5)"

F = "Arial"
H_FILL = PatternFill("solid", fgColor="1F5A42")
H_FONT = Font(name=F, bold=True, color="FFFFFF", size=10)
BODY = Font(name=F, size=10)
INPUT_FILL = PatternFill("solid", fgColor="FFF4C2")   # colunas que a equipe atualiza
TITLE = Font(name=F, bold=True, size=14, color="1F5A42")
BOLD = Font(name=F, bold=True, size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
thin = Side(style="thin", color="D5DBD2")

def cabecalho(ws, cols, larguras, input_cols=(), nota=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font = H_FONT; cell.fill = H_FILL; cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = larguras[i - 1]
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 32
    if nota:
        ws.cell(row=1, column=1).comment = Comment(nota, "Encaminha Goiânia")
    return set(input_cols)

def linha(ws, r, valores, input_cols=(), wrap_cols=()):
    for i, v in enumerate(valores, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = BODY
        c.alignment = WRAP if i in wrap_cols else TOP
        if i in input_cols:
            c.fill = INPUT_FILL

wb = Workbook()

# ---------- Leia-me ----------
ws = wb.active; ws.title = "Leia-me"
ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 100
ws["A1"] = "Banco de dados do Encaminha Goiânia"; ws["A1"].font = TITLE
textos = [
 ("Para que serve", "Guarda os dados do app fora dele: serviços e contatos da rede, textos das políticas, temas do pedido livre e os atendimentos registrados de forma anônima. A equipe gestora atualiza aqui; o app lê daqui."),
 ("Serviços", "Um serviço por linha. Atualize telefone, endereço, horário e observação. Ao conferir, mude a Situação e preencha Conferido em e Conferido por. A coluna Precisa conferir? avisa sozinha: Sim quando nunca foi conferido ou passou de 180 dias."),
 ("Políticas", "Nome, valor, passo a passo e documentos de cada política. Uma linha do passo a passo por linha da célula. Quem tem direito (o critério) fica no código do app e é testado; para mudar um critério, escreva o pedido na coluna Mudança de critério pedida."),
 ("Temas", "Situações que o app reconhece no pedido livre. Palavras-chave sem acento, separadas por vírgula. Uma palavra curta (até 3 letras) precisa aparecer inteira; as demais valem no começo de uma palavra."),
 ("Registros", "Preenchida pelo app quando o profissional toca em Registrar atendimento. Não edite à mão. Só dados anônimos: nada de nome, iniciais, texto livre, renda ou idade exatas."),
 ("Resumo", "Contagens por fórmula sobre a aba Registros. Atualiza sozinha quando chegam registros."),
 ("Fontes", "De onde veio cada tabela e quando foi exportada."),
 ("Como editar", "Não mude a coluna ID nem a ordem das colunas: o app usa os IDs para ligar políticas, temas e serviços. Para incluir um serviço novo, use uma linha nova no fim, com um ID novo, curto e sem espaço (ex.: caps-ad-leste)."),
 ("Legenda", "Células em amarelo claro: colunas de conferência que a equipe preenche."),
 ("LGPD", "Nunca registre nome, CPF, NIS, endereço, telefone ou e-mail de pessoa atendida. Raça, religião, orientação sexual e dados de saúde de uma pessoa identificável são dados sensíveis. Os registros do app são anônimos e situações sensíveis com 1 ou 2 casos não devem ser divulgadas."),
 ("Acesso", "Edição só para a equipe gestora. Para profissionais que só consultam, compartilhe como Leitor."),
]
for i, (a, b) in enumerate(textos, 3):
    ws.cell(row=i, column=1, value=a).font = BOLD
    c = ws.cell(row=i, column=2, value=b); c.font = BODY; c.alignment = WRAP
    ws.cell(row=i, column=1).alignment = TOP
ws.cell(row=3 + len(textos) - 3, column=2).fill = INPUT_FILL  # Legenda

# ---------- Serviços ----------
ws = wb.create_sheet("Serviços")
cols = ["ID (não altere)", "Nome", "Tipo", "Município", "Região de Goiânia", "Endereço", "Telefones (separe com /)", "E-mail", "Horário", "Site", "Observação", "Fonte", "Situação", "Conferido em", "Conferido por (setor)", "Precisa conferir?"]
larg = [18, 42, 22, 20, 16, 44, 34, 30, 24, 34, 60, 34, 13, 13, 18, 13]
ic = {13, 14, 15}
S = d["S"]
cabecalho(ws, cols, larg, ic, f"Fonte: {ORIGEM}, {len(S)} linhas, exportado em {HOJE}. Exportado por Claude. Detalhes na aba Fontes.")
for r, s in enumerate(S, 2):
    obs = s.get("obs", "") or ""
    m = re.search(r"Fonte:\s*([^.]+(?:\.[^ ][^.]*)?)", obs)
    fonte = m.group(1).strip() if m else "Protótipo do app (ver aba Fontes)"
    conferir = bool(re.search(r"confirm|conferid|conferir", obs, re.I))
    dm = re.search(r"consultado em (\d{2}/\d{2}/\d{4})", obs) or re.search(r"Fonte:[^.]*?(\d{2}/\d{2}/\d{4})", obs)
    linha(ws, r, [s["id"], s.get("nome", ""), s.get("tipo", ""), s.get("mun", ""), s.get("reg", ""), s.get("end", ""), " / ".join(s.get("tel", []) or []),
                  s.get("email", ""), s.get("hor", ""), s.get("site", ""), obs, fonte, "A conferir" if conferir else "Ativo",
                  dm.group(1) if dm else None, ("Levantamento de 08/10/2026" if "consultado" in dm.group(0) else "Data da fonte") if dm else None, None], ic, {2, 6, 11, 12})
    if dm:
        ws.cell(row=r, column=14).number_format = "DD/MM/YYYY"
        dd, mm, aa = dm.group(1).split("/")
        import datetime
        ws.cell(row=r, column=14).value = datetime.date(int(aa), int(mm), int(dd))
ULT = len(S) + 1 + 200
for r in range(2, ULT + 1):
    c = ws.cell(row=r, column=16, value=f'=IF($A{r}="","",IF($M{r}="Inativo","Não",IF($N{r}="","Sim",IF(TODAY()-$N{r}>180,"Sim","Não"))))')
    c.font = BODY; c.alignment = TOP
    if r > len(S) + 1:
        for col in ic: ws.cell(row=r, column=col).fill = INPUT_FILL
        ws.cell(row=r, column=14).number_format = "DD/MM/YYYY"
dv = DataValidation(type="list", formula1='"Ativo,A conferir,Inativo"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"M2:M{ULT}")
ws.auto_filter.ref = f"A1:P{len(S)+1}"

# ---------- Políticas ----------
ws = wb.create_sheet("Políticas")
cols = ["ID (não altere)", "Nome", "Nome curto", "Esfera", "Valor ou benefício", "Passo a passo (um passo por linha)", "Documentos (um por linha)", "Onde ir (tipo de lugar)", "Outros contatos (IDs da aba Serviços)", "Fonte", "Conferido em", "Mudança de critério pedida"]
larg = [14, 40, 22, 12, 28, 90, 44, 16, 30, 28, 13, 36]
ic = {11, 12}
P = d["P"]
cabecalho(ws, cols, larg, ic, f"Fonte: {ORIGEM}, {len(P)} linhas, exportado em {HOJE}. Exportado por Claude. Detalhes na aba Fontes.")
ESF = {"nac": "Nacional", "est": "Estadual", "mun": "Municipal", "rede": "Rede de proteção"}
for r, p in enumerate(P, 2):
    linha(ws, r, [p["id"], p["nome"], p["curto"], ESF.get(p["esf"], p["esf"]), p["val"], "\n".join(p["flow"]), "\n".join(p["docs"]), p["where"], ", ".join(p["extra"]), "Ver aba Fontes e docs/levantamento-2026-10-08.md", None, None], ic, {2, 5, 6, 7, 12})
    ws.cell(row=r, column=11).number_format = "DD/MM/YYYY"
dv = DataValidation(type="list", formula1='"Nacional,Estadual,Municipal,Rede de proteção"', allow_blank=False)
ws.add_data_validation(dv); dv.add(f"D2:D{len(P)+50}")

# ---------- Temas ----------
ws = wb.create_sheet("Temas")
cols = ["Chave (não altere)", "Título", "Palavras-chave (sem acento, separadas por vírgula)", "Liga a situação do app", "Passo a passo (um por linha)", "Contatos (IDs da aba Serviços)", "Frase de urgência"]
larg = [16, 40, 60, 18, 80, 34, 50]
T = d["T"]
cabecalho(ws, cols, larg, (), f"Fonte: {ORIGEM}, {len(T)} linhas, exportado em {HOJE}. Exportado por Claude. Detalhes na aba Fontes.")
for r, t in enumerate(T, 2):
    linha(ws, r, [t["k"], t["t"], ", ".join(t["palavras"]), ", ".join(t["on"]), "\n".join(t["passos"]), ", ".join(t["servicos"]), re.sub(r"<[^>]+>", "", t["urg"])], (), {2, 3, 5, 7})

# ---------- Registros ----------
ws = wb.create_sheet("Registros")
cols = ["ID", "Dia", "Hora", "Município", "Região de Goiânia", "Gênero", "Faixa de idade", "Pessoas na casa", "Renda por pessoa", "Trabalho", "Moradia", "Escolaridade", "Assuntos", "Situações", "Políticas indicadas", "Lugares do roteiro", "Temas do pedido livre", "Pedido sem orientação", "Urgência", "Respostas que faltaram"]
larg = [14, 12, 7, 20, 14, 14, 13, 10, 20, 22, 14, 20, 40, 50, 50, 40, 30, 12, 10, 22]
cabecalho(ws, cols, larg, (), "Preenchida pelo app (botão Registrar atendimento). Não edite à mão. Só dados anônimos; listas separadas por ponto e vírgula.")
ws.freeze_panes = "A2"

# ---------- Resumo ----------
ws = wb.create_sheet("Resumo")
ws.column_dimensions["A"].width = 44; ws.column_dimensions["B"].width = 14; ws.column_dimensions["C"].width = 12
ws["A1"] = "Resumo dos atendimentos registrados"; ws["A1"].font = TITLE
ws["A2"] = "Contagens por fórmula sobre a aba Registros. Listas separadas por ponto e vírgula são contadas pelo nome."; ws["A2"].font = BODY
def bloco(r, titulo, itens, col, pct=True):
    ws.cell(row=r, column=1, value=titulo).font = H_FONT; ws.cell(row=r, column=1).fill = H_FILL
    ws.cell(row=r, column=2, value="Atendimentos").font = H_FONT; ws.cell(row=r, column=2).fill = H_FILL
    ws.cell(row=r, column=3, value="% do total").font = H_FONT; ws.cell(row=r, column=3).fill = H_FILL
    for i, it in enumerate(itens, r + 1):
        ws.cell(row=i, column=1, value=it).font = BODY
        f = f'=COUNTIF(Registros!${col}$2:${col},"*"&$A{i}&"*")' if pct == "lista" else f'=COUNTIF(Registros!${col}$2:${col},$A{i})'
        ws.cell(row=i, column=2, value=f).font = BODY
        c = ws.cell(row=i, column=3, value=f"=IF($B$5=0,0,B{i}/$B$5)"); c.font = BODY; c.number_format = "0.0%"
    return r + len(itens) + 2
ws["A4"] = "Indicador"; ws["B4"] = "Valor"
for c in ("A4", "B4"): ws[c].font = H_FONT; ws[c].fill = H_FILL
ind = [("Atendimentos registrados", "=COUNTA(Registros!A2:A)", None),
       ("Com urgência", '=COUNTIF(Registros!S2:S,"Sim")', None),
       ("Com urgência (%)", "=IF(B5=0,0,B6/B5)", "0.0%"),
       ("Renda até R$ 218 por pessoa", '=COUNTIF(Registros!I2:I,"Sem renda")+COUNTIF(Registros!I2:I,"Até R$ 109")+COUNTIF(Registros!I2:I,"De R$ 110 a R$ 218")', None),
       ("Renda até R$ 218 por pessoa (%)", "=IF(B5=0,0,B8/B5)", "0.0%"),
       ("Pedidos livres sem orientação pronta", '=COUNTIF(Registros!R2:R,"Sim")', None),
       ("Serviços que precisam de conferência", "=COUNTIF('Serviços'!P2:P,\"Sim\")", None)]
for i, (a, f, fmt) in enumerate(ind, 5):
    ws.cell(row=i, column=1, value=a).font = BODY
    c = ws.cell(row=i, column=2, value=f); c.font = BOLD
    if fmt: c.number_format = fmt
NEC = ["Renda e benefícios", "Comida", "Moradia", "Saúde e remédios", "Saúde mental e dependências", "Violência e proteção", "Educação", "Trabalho", "Documentos", "Direito negado"]
MUN = ["Goiânia", "Aparecida de Goiânia", "Senador Canedo", "Trindade", "Goianira", "Abadia de Goiás", "Aragoiânia", "Bela Vista de Goiás", "Bonfinópolis", "Brazabrantes", "Caldazinha", "Caturaí", "Goianápolis", "Guapó", "Hidrolândia", "Inhumas", "Nerópolis", "Nova Veneza", "Santo Antônio de Goiás", "Terezópolis de Goiás"]
RENDA = ["Sem renda", "Até R$ 109", "De R$ 110 a R$ 218", "Até 1/4 do salário", "Até 1/2 salário", "Até 1 salário", "Até 3 salários", "Mais de 3 salários", "Não informada"]
IDADE = ["Até 13", "14 a 15", "16 a 17", "18 a 29", "30 a 59", "60 a 64", "65 ou mais", "Não informada"]
r = 13
r = bloco(r, "Assuntos", NEC, "M", "lista")
r = bloco(r, "Renda por pessoa", RENDA, "I")
r = bloco(r, "Faixa de idade", IDADE, "G")
r = bloco(r, "Município", MUN, "D")

# ---------- Fontes ----------
ws = wb.create_sheet("Fontes")
cols = ["Escrito em", "Fonte", "Objeto", "Obtido por", "Pedido", "Parâmetros", "Data", "Pedido por", "Linhas", "Mudanças depois de obter", "Observações"]
larg = [22, 28, 30, 26, 44, 16, 12, 14, 8, 30, 50]
cabecalho(ws, cols, larg)
ws.freeze_panes = "A2"
reg = [
 (f"Serviços!A1:P{len(S)+1}", ORIGEM, "lista S (serviços e contatos)", "Exportado por Claude do código do app", "Criar banco de dados independente no Google Planilhas", "", HOJE, "", len(S), "Fonte e data extraídas da observação quando havia", "Contatos vêm do protótipo e do levantamento de 08/10/2026 (docs/levantamento-2026-10-08.md)."),
 (f"Políticas!A1:L{len(P)+1}", ORIGEM, "lista P (políticas)", "Exportado por Claude do código do app", "Criar banco de dados independente no Google Planilhas", "", HOJE, "", len(P), "Critérios de direito não exportados (ficam no código)", ""),
 (f"Temas!A1:G{len(T)+1}", ORIGEM, "lista TEMAS (pedido livre)", "Exportado por Claude do código do app", "Criar banco de dados independente no Google Planilhas", "", HOJE, "", len(T), "Marcação HTML retirada da frase de urgência", ""),
]
for i, v in enumerate(reg, 2):
    linha(ws, i, list(v), (), {2, 5, 10, 11})

for w in wb.worksheets:
    w.sheet_view.showGridLines = True
wb.save(out)
print("ok", out)
